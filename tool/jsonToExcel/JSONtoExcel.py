from __future__ import annotations

import copy
import json
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


HEADER_ROWS = 5
KEY_MARK = "【KEY】"


def usage() -> None:
    print("Usage:")
    print("  JSONtoExcel.bat input.json [output.xlsx] [template.xlsx]")
    print()
    print("Rows generated for excel-to-json compatibility:")
    print("  Row 1: Chinese/comment name, with 【KEY】 on primary key columns")
    print("  Row 2: Field key")
    print("  Row 3: Type: int / float / string / int[] / any")
    print("  Row 4: server")
    print("  Row 5: client")


def is_record(value: Any) -> bool:
    return isinstance(value, dict) and any(not isinstance(v, dict) for v in value.values())


def flatten_json(data: Any) -> list[OrderedDict[str, Any]]:
    rows: list[OrderedDict[str, Any]] = []

    if isinstance(data, list):
        for item in data:
            if isinstance(item, dict):
                rows.append(OrderedDict(item))
            else:
                rows.append(OrderedDict(value=item))
        return rows

    if not isinstance(data, dict):
        return [OrderedDict(value=data)]

    def walk(node: Any, key_path: list[str]) -> None:
        if isinstance(node, dict) and is_record(node):
            row = OrderedDict()
            for key, value in node.items():
                row[key] = value

            # Most project JSON is keyed by id and also stores id in the row.
            # If id is missing, promote the last outer key to id for excel-to-json.
            if "id" not in row and key_path and key_path[-1].lstrip("-").isdigit():
                row["id"] = int(key_path[-1])
                row.move_to_end("id", last=False)
            elif "id" not in row:
                for index, key in enumerate(key_path, start=1):
                    row[f"_key{index}"] = key
            rows.append(row)
            return

        if isinstance(node, dict):
            for key, value in node.items():
                walk(value, key_path + [str(key)])
            return

        row = OrderedDict()
        for index, key in enumerate(key_path, start=1):
            row[f"_key{index}"] = key
        row["value"] = node
        rows.append(row)

    for key, value in data.items():
        walk(value, [str(key)])
    return rows


def infer_type(values: list[Any]) -> str:
    present = [value for value in values if value not in (None, "")]
    if not present:
        return "string"

    if all(isinstance(value, bool) for value in present):
        return "int"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in present):
        return "int"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present):
        return "float"
    if all(
        isinstance(value, list)
        and all(isinstance(item, int) and not isinstance(item, bool) for item in value)
        for value in present
    ):
        return "int[]"
    if any(isinstance(value, (dict, list)) for value in present):
        return "any"
    return "string"


def ordered_fields(rows: list[OrderedDict[str, Any]]) -> list[str]:
    fields: list[str] = []
    for preferred in ("id", "_key1", "_key2", "_key3", "name"):
        if any(preferred in row for row in rows) and preferred not in fields:
            fields.append(preferred)

    for row in rows:
        for key in row.keys():
            if key not in fields:
                fields.append(key)
    return fields


def serialize_value(value: Any, field_type: str) -> Any:
    if value is None:
        return ""
    if field_type == "int[]":
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if field_type == "any" and isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, bool):
        return 1 if value else 0
    return value


def copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)


def load_template_headers(template_path: Path) -> tuple[list[str], dict[str, list[Any]], Any | None]:
    if not template_path:
        return [], {}, None

    template_wb = load_workbook(template_path)
    template_ws = template_wb.worksheets[0]
    fields: list[str] = []
    header_by_field: dict[str, list[Any]] = {}

    for col in range(1, template_ws.max_column + 1):
        field = template_ws.cell(2, col).value
        if field in (None, ""):
            continue
        field = str(field)
        fields.append(field)
        header_by_field[field] = [template_ws.cell(row, col).value for row in range(1, HEADER_ROWS + 1)]

    return fields, header_by_field, template_ws


def build_workbook(json_path: Path, output_path: Path, template_path: Path | None) -> None:
    with open(json_path, "r", encoding="utf-8") as file:
        data = json.load(file, object_pairs_hook=OrderedDict)

    rows = flatten_json(data)
    if not rows:
        raise ValueError("No rows found in JSON.")

    template_fields, template_headers, template_ws = load_template_headers(template_path) if template_path else ([], {}, None)
    inferred_fields = ordered_fields(rows)
    fields = [field for field in template_fields if field in inferred_fields]
    fields.extend(field for field in inferred_fields if field not in fields)

    type_by_field = {
        field: infer_type([row.get(field) for row in rows])
        for field in fields
    }

    wb = Workbook()
    ws = wb.active
    ws.title = json_path.stem[:31] or "JSON"

    key_fields = [field for field in fields if field == "id"]
    if not key_fields:
        key_fields = [field for field in fields if field.startswith("_key")]
    if not key_fields and fields:
        key_fields = [fields[0]]

    for col, field in enumerate(fields, start=1):
        headers = template_headers.get(field)
        if headers:
            values = headers[:]
            if values[2] in (None, ""):
                values[2] = type_by_field[field]
            if values[3] in (None, ""):
                values[3] = "server"
            if values[4] in (None, ""):
                values[4] = "client"
        else:
            display_name = field + (KEY_MARK if field in key_fields else "")
            values = [display_name, field, type_by_field[field], "server", "client"]

        for row_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col)
            cell.value = value
            if template_ws is not None and field in template_headers:
                template_col = template_fields.index(field) + 1
                copy_cell_style(template_ws.cell(row_index, template_col), cell)

    for row_index, row_data in enumerate(rows, start=HEADER_ROWS + 1):
        for col, field in enumerate(fields, start=1):
            value = serialize_value(row_data.get(field), type_by_field[field])
            ws.cell(row_index, col).value = value

    if template_ws is not None:
        for col in range(1, len(fields) + 1):
            letter = get_column_letter(col)
            template_col = template_fields.index(fields[col - 1]) + 1 if fields[col - 1] in template_fields else None
            if template_col:
                template_letter = get_column_letter(template_col)
                ws.column_dimensions[letter].width = template_ws.column_dimensions[template_letter].width
        for row in range(1, HEADER_ROWS + 1):
            ws.row_dimensions[row].height = template_ws.row_dimensions[row].height
    else:
        for col, field in enumerate(fields, start=1):
            max_len = max(len(str(field)), *(len(str(row.get(field, ""))) for row in rows[:200]))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 50)

    ws.freeze_panes = "A6"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def resolve_output_path(json_path: Path, output_arg: Path | None) -> Path:
    if output_arg is None:
        return json_path.with_suffix(".xlsx")
    if output_arg.suffix.lower() == ".xlsx":
        return output_arg
    return output_arg / f"{json_path.stem}.xlsx"


def resolve_template_path(json_path: Path, template_arg: Path | None) -> Path | None:
    if template_arg is None:
        return None
    if template_arg.is_dir():
        same_name = template_arg / f"{json_path.stem}.xlsx"
        return same_name if same_name.exists() else None
    return template_arg if template_arg.exists() else None


def convert_one(json_path: Path, output_arg: Path | None, template_arg: Path | None) -> Path:
    output_path = resolve_output_path(json_path, output_arg)
    template_path = resolve_template_path(json_path, template_arg)
    build_workbook(json_path, output_path, template_path)
    return output_path


def convert_folder(input_dir: Path, output_dir: Path | None, template_arg: Path | None) -> list[Path]:
    json_files = sorted(input_dir.glob("*.json"))
    if not json_files:
        raise ValueError(f"No .json files found in folder: {input_dir}")

    output_paths: list[Path] = []
    for json_path in json_files:
        output_paths.append(convert_one(json_path, output_dir, template_arg))
    return output_paths


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        usage()
        return 1

    input_path = Path(argv[1]).resolve()
    if not input_path.exists():
        print(f"Input not found: {input_path}")
        return 1

    output_arg = Path(argv[2]).resolve() if len(argv) >= 3 and argv[2] else None
    template_path = Path(argv[3]).resolve() if len(argv) >= 4 else None
    if template_path and not template_path.exists():
        print(f"Template Excel/folder not found: {template_path}")
        return 1

    if input_path.is_dir():
        output_paths = convert_folder(input_path, output_arg, template_path)
        print(f"Converted {len(output_paths)} JSON files.")
        for output_path in output_paths:
            print(f"Saved: {output_path}")
    else:
        output_path = convert_one(input_path, output_arg, template_path)
        print(f"Saved: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
